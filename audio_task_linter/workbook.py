"""openpyxl helpers: load a workbook twice (formulas + cached values) and index it."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

import openpyxl
from openpyxl.utils import get_column_letter

from .numbers import ERROR_VALUES, Num, within

_UNIT_SCALES = (1.0, 1000.0, 1_000_000.0, 0.001, 0.000001)


@dataclass
class Cell:
    sheet: str
    coord: str
    row: int
    col: int
    value: object            # cached value (data_only)
    formula: Optional[str]   # formula string if the cell is a formula, else None

    @property
    def is_formula(self) -> bool:
        return self.formula is not None

    @property
    def is_number(self) -> bool:
        return isinstance(self.value, (int, float)) and not isinstance(self.value, bool)

    @property
    def ref(self) -> str:
        return f"{self.sheet}!{self.coord}"


@dataclass
class Workbook:
    path: Path
    cells: dict = field(default_factory=dict)        # (sheet, coord) -> Cell
    sheets: list = field(default_factory=list)
    hidden_sheets: list = field(default_factory=list)
    external_links: list = field(default_factory=list)
    defined_names: dict = field(default_factory=dict)
    comments: list = field(default_factory=list)
    images: list = field(default_factory=list)      # sheet names carrying embedded images
    row_labels: dict = field(default_factory=dict)   # (sheet, row) -> label text
    col_headers: dict = field(default_factory=dict)  # (sheet, col) -> header text
    sha256: str = ""

    def numeric_cells(self) -> Iterable[Cell]:
        return (c for c in self.cells.values() if c.is_number)

    def formula_cells(self) -> Iterable[Cell]:
        return (c for c in self.cells.values() if c.is_formula)

    def literal_numbers(self) -> Iterable[Cell]:
        return (c for c in self.cells.values() if c.is_number and not c.is_formula)

    def error_cells(self) -> list[Cell]:
        return [c for c in self.cells.values()
                if isinstance(c.value, str) and c.value.strip().upper() in ERROR_VALUES]

    def sheet_names_lower(self) -> dict:
        return {s.lower(): s for s in self.sheets}

    def label_for(self, cell: Cell) -> str:
        return self.row_labels.get((cell.sheet, cell.row), "")

    def header_for(self, cell: Cell) -> str:
        return self.col_headers.get((cell.sheet, cell.col), "")

    def find_value(self, target: Num, tol: Optional[dict] = None, sheet: Optional[str] = None,
                   allow_scales: bool = True) -> list[tuple[Cell, float]]:
        """Cells whose numeric value equals target (any display-unit scaling). Returns (cell, scale)."""
        scales = list(_UNIT_SCALES) if (allow_scales and target.unit == "" and abs(target.value) >= 1) else [1.0]
        variants = []
        for s in scales:
            variants.append((target.value * s, s))
        if target.is_percent:
            # 10.8% is stored as 0.108 (scale 0.01) or as 10.8
            variants.append((target.value / 100.0, 0.01))
        if target.unit == "pp":
            variants.append((target.value / 100.0, 0.01))
        hits = []
        for c in self.numeric_cells():
            if sheet and c.sheet.lower() != sheet.lower():
                continue
            v = float(c.value)
            for tv, s in variants:
                if within(v, tv, _scaled_tol(tol, s)):
                    hits.append((c, s))
                    break
        return hits


def _scaled_tol(tol: Optional[dict], scale: float) -> Optional[dict]:
    if tol is None or tol["kind"] == "relative":
        return tol
    t = dict(tol)
    if tol["kind"] == "absolute_pp":
        # ±0.5 percentage points on a value stored as 0.108 => ±0.005
        t["kind"] = "absolute"
        t["value"] = tol["value"] * (0.01 if scale == 0.01 else 1.0)
    else:
        t["value"] = tol["value"] * scale
    return t


def load_workbook(path: Path, max_cells_per_sheet: int = 400_000) -> Workbook:
    path = Path(path)
    wb = Workbook(path=path)
    wb.sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    wf = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    wv = openpyxl.load_workbook(path, data_only=True, keep_links=True)
    wb.sheets = list(wf.sheetnames)
    wb.hidden_sheets = [ws.title for ws in wf.worksheets if ws.sheet_state != "visible"]
    try:
        wb.external_links = [getattr(l, "file_link", None) and l.file_link.Target for l in wf._external_links]
    except Exception:  # pragma: no cover
        wb.external_links = []
    try:
        wb.defined_names = {n: d.attr_text for n, d in wf.defined_names.items()}
    except Exception:  # pragma: no cover
        wb.defined_names = {}
    for ws_f in wf.worksheets:
        ws_v = wv[ws_f.title]
        if getattr(ws_f, '_images', None):
            wb.images.append(ws_f.title)
        n = 0
        for row in ws_f.iter_rows():
            for cf in row:
                if cf.value is None:
                    continue
                n += 1
                if n > max_cells_per_sheet:
                    break
                cv = ws_v[cf.coordinate].value
                formula = None
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    formula = cf.value
                elif cf.data_type == "f":
                    formula = str(cf.value)
                cell = Cell(ws_f.title, cf.coordinate, cf.row, cf.column, cv, formula)
                wb.cells[(ws_f.title, cf.coordinate)] = cell
                if cf.comment is not None:
                    wb.comments.append(cell.ref)
                if isinstance(cf.value, str) and formula is None:
                    key = (ws_f.title, cf.row)
                    if key not in wb.row_labels:
                        wb.row_labels[key] = cf.value.strip()
                    if cf.row <= 12 and (ws_f.title, cf.column) not in wb.col_headers:
                        wb.col_headers[(ws_f.title, cf.column)] = cf.value.strip()
        # date-like header rows: also record numeric headers (years / periods) in the top rows
        for row in ws_v.iter_rows(min_row=1, max_row=12):
            for cv in row:
                key = (ws_f.title, cv.column)
                existing = wb.col_headers.get(key, "")
                has_year = bool(re.search(r"(?:19|20)\d{2}|(?:FY|CY|')\s?\d{2}\b", existing))
                if isinstance(cv.value, (int, float)) and not isinstance(cv.value, bool) and 1990 <= cv.value <= 2100:
                    if not has_year:
                        wb.col_headers[key] = (existing + " " + str(int(cv.value))).strip()
                elif hasattr(cv.value, "year"):
                    if not has_year:
                        wb.col_headers[key] = (existing + " " + cv.value.strftime("%Y-%m-%d")).strip()
    return wb


def words(s: str) -> set[str]:
    stop = {"the", "a", "an", "of", "for", "on", "in", "to", "and", "or", "is", "does", "do", "sheet", "tab",
            "at", "as", "by", "with", "within", "case", "base", "new", "total", "per", "from", "value", "values"}
    return {w for w in re.findall(r"[a-z0-9%]+", s.lower()) if w not in stop and len(w) > 1}


def label_score(text: str, label: str, header: str = "") -> float:
    """Word-overlap score between a rubric phrase and a row label (+ column header)."""
    tw = words(text)
    if not tw:
        return 0.0
    lw = words(label) | words(header)
    return len(tw & lw) / len(tw)


def coord(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"
