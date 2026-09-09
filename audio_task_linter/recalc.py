"""Recalculate a workbook after a perturbation.

Primary engine: LibreOffice Calc headless (faithful to Excel semantics, ~1s per run).
Fallback engine: the pure-Python `formulas` package (pip install formulas) when soffice is absent.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Optional

import openpyxl


def soffice_path() -> Optional[str]:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    return None


def formulas_available() -> bool:
    try:
        import formulas  # noqa: F401
        return True
    except ImportError:
        return False


def engine() -> Optional[str]:
    if soffice_path():
        return "libreoffice"
    if formulas_available():
        return "formulas"
    return None


def recalculate(src: Path, out_dir: Path, timeout: int = 180) -> Path:
    """Convert src -> xlsx via LibreOffice; formulas without cached values are recomputed on load."""
    exe = soffice_path()
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not found on PATH; cannot recalculate")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-", dir=out_dir))
    cmd = [exe, "--headless", "--norestore", "--nologo",
           f"-env:UserInstallation=file://{profile}",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir), str(src)]
    env = dict(os.environ, HOME=str(profile))
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=env)
    target = out_dir / (Path(src).stem + ".xlsx")
    if proc.returncode != 0 or not target.exists():
        msg = (proc.stderr + proc.stdout).strip()[:400]
        if "could not be loaded" in msg:
            msg += " (is libreoffice-calc installed? core-only installs cannot open spreadsheets)"
        raise RuntimeError(f"LibreOffice recalculation failed: rc={proc.returncode} {msg}")
    return target


def perturb_and_recalc(gold_path: Path, edits: dict, work_dir: Path) -> Path:
    """Apply {(sheet, coord): new_value} edits to a copy of the gold and return a recalculated copy.

    openpyxl drops cached values on save, so every formula is recomputed by the engine.
    """
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    eng = engine()
    if eng == "libreoffice":
        wb = openpyxl.load_workbook(gold_path)  # formulas preserved
        for (sheet, coord), val in edits.items():
            wb[sheet][coord].value = val
        tmp = work_dir / "perturbed_input.xlsx"
        wb.save(tmp)
        return recalculate(tmp, work_dir / "recalc")
    if eng == "formulas":
        return _perturb_with_formulas(gold_path, edits, work_dir)
    raise RuntimeError("no recalculation engine: install libreoffice-calc or `pip install formulas`")


def _perturb_with_formulas(gold_path: Path, edits: dict, work_dir: Path) -> Path:
    import formulas  # type: ignore

    gold_path = Path(gold_path)
    xl = formulas.ExcelModel().loads(str(gold_path)).finish()
    book = f"[{gold_path.name}]".upper()
    inputs = {f"'{book}{sheet.upper()}'!{coord.upper()}": val for (sheet, coord), val in edits.items()}
    sol = xl.calculate(inputs=inputs)
    wb = openpyxl.load_workbook(gold_path)
    for (sheet, coord), val in edits.items():
        wb[sheet][coord].value = val
    # write computed values back over the formulas so the result reads like a cached workbook
    wf = openpyxl.load_workbook(gold_path)
    for key, cell in sol.items():
        if "!" not in key:
            continue
        sheet_part, coord = key.rsplit("!", 1)
        sheet = sheet_part.split("]")[-1].strip("'")
        ws = _sheet_ci(wb, sheet)
        if ws is None or ":" in coord:
            continue
        src = _sheet_ci(wf, sheet)[coord]
        if isinstance(src.value, str) and src.value.startswith("="):
            v = cell.value[0][0] if hasattr(cell, "value") else cell
            try:
                v = v.item() if hasattr(v, "item") else v
            except Exception:
                pass
            ws[coord].value = None if v is formulas.functions.EMPTY else v  # type: ignore[attr-defined]
    out = work_dir / "perturbed_recalc_formulas.xlsx"
    wb.save(out)
    return out


def _sheet_ci(wb, name: str):
    for ws in wb.worksheets:
        if ws.title.upper() == name.upper():
            return ws
    return None
