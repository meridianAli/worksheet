import pytest

from audio_task_linter import recalc
from audio_task_linter.__main__ import lint_dir
from fixtures import build_bundle


def rule_ids(report, severity=None):
    return {f.rule for f in report.findings if severity is None or f.severity == severity}


def test_good_bundle_passes(tmp_path):
    root = build_bundle(tmp_path / "good", good=True)
    rep = lint_dir(root, recalc=True, work_dir=tmp_path / "work")
    assert rep.ok, [f.__dict__ for f in rep.findings if f.severity == "error"]
    # fixtures are written by openpyxl, so the provenance rule legitimately fires on them
    assert rule_ids(rep, "warning") <= {"V001"}


def test_bad_bundle_flags_every_planted_defect(tmp_path):
    root = build_bundle(tmp_path / "bad", good=False, hints=True)
    rep = lint_dir(root, recalc=recalc.engine() is not None, work_dir=tmp_path / "work")
    assert not rep.ok
    errors = rule_ids(rep, "error")
    expected_errors = {"R001", "R002", "R003", "R004", "R006", "R008", "R009", "S004", "G005", "I001", "G008", "G009"}
    assert expected_errors <= errors, expected_errors - errors
    warnings = rule_ids(rep, "warning")
    expected_warnings = {"R005", "R010", "R011", "R012", "R013", "R014", "I002"}
    assert expected_warnings <= warnings, expected_warnings - warnings


@pytest.mark.skipif(recalc.engine() is None, reason="no recalculation engine (libreoffice-calc or `formulas`)")
def test_perturbations_are_recalculated(tmp_path):
    root = build_bundle(tmp_path / "bad", good=False, hints=True)
    rep = lint_dir(root, recalc=True, work_dir=tmp_path / "work")
    p001_errors = [f for f in rep.findings if f.rule == "P001" and f.severity == "error"]
    # P-1 wrong target, P-2 no-op perturbation, P-4 wrong target
    assert {f.location for f in p001_errors} == {"PT-1", "PT-2", "PT-4"}
    assert any(f.rule == "P002" and f.location == "PT-5" for f in rep.findings)


def test_missing_pieces_reported(tmp_path):
    root = tmp_path / "empty"
    root.mkdir()
    (root / "rubric.md").write_text("Output Validation\n- Does X equal 5 (+/-2%)? +3\n")
    rep = lint_dir(root, recalc=False, work_dir=tmp_path / "work")
    x001 = [f for f in rep.findings if f.rule == "X001" and f.severity == "error"]
    assert x001 and "audio recording" in x001[0].message and "gold output workbook" in x001[0].message
    assert any(f.rule == "S001" for f in rep.findings)


def test_json_output_roundtrip(tmp_path):
    root = build_bundle(tmp_path / "good", good=True)
    rep = lint_dir(root, recalc=False, work_dir=tmp_path / "work")
    d = rep.to_dict()
    assert d["summary"]["error"] == 0
    assert d["bundle"]["gold_workbook"] == "Meridian-abc123-gold-output.xlsx"


def test_hidden_hardcode_messages(tmp_path):
    root = build_bundle(tmp_path / "bad", good=False, hints=True)
    rep = lint_dir(root, recalc=False, work_dir=tmp_path / "work")
    g008 = [f for f in rep.findings if f.rule == "G008"]
    assert any("1.075" in f.message and f.severity == "error" for f in g008)      # two cells, not spoken, not in input
    assert any("8.5" in f.message and f.severity == "warning" for f in g008)      # one cell
    assert not any("0.05" in f.message and f.severity != "info" for f in g008)    # spoken as "five points"


def test_spoken_numbers():
    from audio_task_linter.spoken import words_to_numbers
    assert words_to_numbers("Four hundred at close, five percent amort, thirteen and a half percent of revenue, two point two five times") == [400, 5, 13.5, 2.25]


def test_duplicate_prompt_across_tasks(tmp_path):
    from audio_task_linter.rules.provenance_rules import check_duplicate_prompts
    a = build_bundle(tmp_path / "a", good=True)
    b = build_bundle(tmp_path / "b", good=True)
    scripts = {}
    reps = [lint_dir(a, recalc=False, work_dir=tmp_path / "w1", scripts_out=scripts),
            lint_dir(b, recalc=False, work_dir=tmp_path / "w2", scripts_out=scripts)]
    check_duplicate_prompts(reps, scripts)
    assert all(any(f.rule == "V005" and f.severity == "error" for f in r.findings) for r in reps)


def test_provenance_and_pii(tmp_path):
    import openpyxl
    root = build_bundle(tmp_path / "p", good=True)
    wb = openpyxl.load_workbook(root / "Meridian-abc123-gold-output.xlsx")
    ws = wb.create_sheet("Sheet1")
    ws["A1"] = "Downloaded from Wall Street Prep - www.wallstreetprep.com"
    ws["A2"] = "Contact: jane.doe@acmeholdings.com"
    ws["A3"] = "Prepared for Zephyr Robotics Inc."
    wb.properties.creator = "Jane Doe"
    wb.save(root / "Meridian-abc123-gold-output.xlsx")
    rep = lint_dir(root, recalc=False, work_dir=tmp_path / "w")
    msgs = " ".join(f.rule + ":" + f.message for f in rep.findings)
    assert "V002" in msgs and "Wall Street Prep" in msgs
    assert "V003" in msgs and "jane.doe@acmeholdings.com" in msgs and "Zephyr Robotics Inc" in msgs
    assert "Jane Doe" not in msgs  # author names in document properties are deliberately ignored
