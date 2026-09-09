"""Build a small but realistic audio-task bundle for tests and demos."""
from __future__ import annotations

from pathlib import Path

import openpyxl

YEARS = [2024, 2025, 2026, 2027, 2028]

SCRIPT = """Rich (MD): Chad, quick debrief from the Vision call. Here's what we need before tomorrow.
Mike (VP): First, on the Inputs tab turn the case selector into a live driver. Case one is Base, two is Bull,
three is Bear. Bull is five points of growth and three points of margin above Base, Bear the same below. Flex them off Base formulaically, don't type the numbers in.
Rich: Tax rate stays at twenty-one percent in Base, NWC at thirteen and a half percent of revenue.
Mike: Then build the model out. Revenue grows at the case growth rate off the twenty twenty-four actual,
EBITDA at the case margin, cash taxes off EBITDA at the tax rate, NWC as a percent of revenue, and free cash
flow available for debt service is EBITDA less taxes less the change in working capital.
Rich: Layer in the new term loan. Four hundred at close, SOFR plus four hundred in Base off the SOFR file we sent,
five percent mandatory amort, sweep everything left over against the balance. Net debt is the ending balance
less cash. Formulas everywhere, I want to be able to flip the case and watch it flow.
Mike: And put a Returns section on the Model tab with the sponsor IRR and MOIC off the equity check
and the exit at the twenty twenty-eight multiple. Add a check row that FCF ties to the debt paydown plus
change in cash.
Rich: Blue font on the hardcoded assumptions, one decimal on the percentages. That's it, thanks Chad.
""" * 2

RUBRIC_GOOD = """Output Validation
- Does Total Revenue for 2028 on the Model sheet equal $1,464.1 in the sheet's displayed units (within +/-2%)? +5
- Does EBITDA for 2028 on the Model sheet equal $439.2 in the sheet's displayed units (within +/-2%)? +4
- Does Free Cash Flow Available for Debt Service for 2028 on the Model sheet equal $329.0 (within +/-2%)? +5
- Does the Term Loan Ending Balance for 2028 on the Model sheet equal $0.0 (within +/- 1.0)? +5
- Does Net Debt for 2028 on the Model sheet equal ($984.3) in the sheet's displayed units (within +/-2%)? +5
- Does the FCF tie-out check row on the Model sheet equal 0 (within +/- 0.5) for 2028? +3

Perturbation
- If Case Number on the Inputs sheet is changed from 1 to 3, does Total Revenue on the Model sheet for 2028 update to $1,215.5 (within +/-2%)? +5
- If the Base Case Tax Rate on the Inputs sheet is changed from 21.0% to 42.0%, do Cash Taxes on the Model sheet for 2028 update to ($184.5) (within +/-2%)? +5
- If the Base Case NWC % Revenue on the Inputs sheet is changed from 13.5% to 27.0%, does Free Cash Flow Available for Debt Service on the Model sheet for 2028 update to $311.1 (within +/-2%)? +5

Presentation
- On the Inputs sheet, do the hardcoded case selector and scenario assumptions use blue font? +2
- On the Model sheet, are the growth and margin lines formatted as percentages with one decimal place? +2

Pitfalls
- Are there any error values (#REF!, #DIV/0!, #VALUE!, #NAME?, #N/A, #NUM!) on the Inputs or Model sheets? -10
- Is any label text truncated due to insufficient column widths on the Model sheet? -5
"""

RUBRIC_BAD = """Output Validation
- Does cell D42 on the Model sheet equal $1,464.1 (within +/-2%)? +5
- Does EBITDA for 2028 equal $439.2 (within +/-2%) and does Net Debt equal ($637.8) (within +/-2%) and does FCF equal $277.3? +12
- Does Total Revenue for 2028 on the Model sheet equal $1,464.1 in the sheet's displayed units (within +/-2%)? +5
- Does Total Revenue for 2028 on the Model sheet equal $1,464.1 in the sheet's displayed units (within +/-2%)? +5
- Is the revenue build correct? +3
- The Returns section should show IRR and MOIC clearly. +4
- Does the SOFR lookup use an XLOOKUP against the SOFR tab? +2
- Does the Base Case Tax Rate on the Inputs sheet equal 21.0% (within +/-2%)? +3
- Does Net Debt for 2028 equal $999.9 (within +/-2%)?
- Does the tie-out check row equal 0 (within +/-2%) for 2028? +3

Perturbation
- If Case Number on the Inputs sheet is changed from 1 to 3, does Total Revenue on the Model sheet for 2028 update to $1,157.6 (within +/-2%) (+/-2%)? +5
- If the Base Case Tax Rate is changed from 21.0% to 21.0%, do Cash Taxes update to ($184.5) (within +/-2%)? +5
- Does the model update when the case changes? +5
- If the Base Case Tax Rate on the Inputs sheet is changed from 21.0% to 42.0%, do Cash Taxes on the Model sheet for 2028 update to ($999.0) (within +/-2%)? +5
- If Exit Multiple on the Inputs sheet is changed from 8.0x to 10.0x, does Total Revenue on the Model sheet for 2028 update to $1,464.1 (within +/-2%)? +5

Presentation

Pitfalls
- Are there any error values (#REF!, #DIV/0!, #VALUE!) on the Model sheet? +10
- On the Transaction sheet, is any label truncated? -5
"""


def _blue(cell):
    cell.font = openpyxl.styles.Font(color="FF0000FF")


def build_gold(path: Path, hidden_hardcode: bool = False):
    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    rows = [
        ("Case Number", 1),
        ("", None),
        ("Assumption", "Base", "Bull", "Bear"),
        ("Revenue Growth", 0.10, "=B4+0.05", "=B4-0.05"),
        ("EBITDA Margin", 0.30, "=B5+0.03", "=B5-0.03"),
        ("Tax Rate", 0.21, "=B6", "=B6"),
        ("NWC % Revenue", 0.135, "=B7", "=B7"),
        ("", None),
        ("Live Revenue Growth", "=CHOOSE($B$1,B4,C4,D4)"),
        ("Live EBITDA Margin", "=CHOOSE($B$1,B5,C5,D5)"),
        ("Live Tax Rate", "=CHOOSE($B$1,B6,C6,D6)"),
        ("Live NWC %", "=CHOOSE($B$1,B7,C7,D7)"),
        ("", None),
        ("Term Loan at Close", 400.0),
        ("SOFR Spread", 0.04),
        ("Mandatory Amort %", 0.05),
        ("Exit Multiple", 8.0),
        ("Equity Check", 300.0),
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is not None and v != "":
                inp.cell(row=r, column=c, value=v)
    for coord in ("B1", "B4", "B5", "B6", "B7", "B14", "B15", "B16", "B17", "B18"):
        _blue(inp[coord])
    for r in (4, 5, 6, 7, 9, 10, 11, 12, 15, 16):
        for col in "BCD":
            inp[f"{col}{r}"].number_format = "0.0%"

    sofr = wb.create_sheet("SOFR")
    sofr["A1"], sofr["B1"] = "Year", "SOFR"
    for i, y in enumerate(YEARS, start=2):
        sofr.cell(row=i, column=1, value=y)
        sofr.cell(row=i, column=2, value=0.05 - 0.005 * (i - 2))

    m = wb.create_sheet("Model")
    m["A1"] = "($ in millions)"
    for i, y in enumerate(YEARS):
        m.cell(row=2, column=2 + i, value=y)
    labels = ["Total Revenue", "% Growth", "EBITDA", "% Margin", "Cash Taxes", "Net Working Capital", "Change in NWC",
              "Free Cash Flow Available for Debt Service", "", "Term Loan Beginning Balance", "Mandatory Amortization",
              "Cash Sweep", "Term Loan Ending Balance", "Interest Expense", "", "Cash Balance", "Net Debt", "",
              "FCF tie-out check", "", "Returns", "Exit Equity Value", "MOIC", "IRR"]
    for i, lab in enumerate(labels, start=3):
        if lab:
            m.cell(row=i, column=1, value=lab)
    cols = ["B", "C", "D", "E", "F"]
    m["B3"] = 1000.0
    _blue(m["B3"])
    for j, col in enumerate(cols):
        prev = cols[j - 1] if j else None
        if prev:
            m[f"{col}3"] = f"={prev}3*(1+Inputs!$B$9)"
            m[f"{col}4"] = f"={col}3/{prev}3-1"
            m[f"{col}4"].number_format = "0.0%"
        m[f"{col}5"] = f"={col}3*Inputs!$B$10"
        m[f"{col}6"] = f"={col}5/{col}3"
        m[f"{col}6"].number_format = "0.0%"
        m[f"{col}7"] = f"=-{col}5*Inputs!$B$11"
        m[f"{col}8"] = f"={col}3*Inputs!$B$12"
        m[f"{col}9"] = f"=-({col}8-{prev}8)" if prev else 0
        m[f"{col}10"] = f"={col}5+{col}7+{col}9"
        # debt
        m[f"{col}12"] = "=Inputs!$B$14" if not prev else f"={prev}15"
        m[f"{col}13"] = f"=-MIN({col}12,Inputs!$B$14*Inputs!$B$16)"
        m[f"{col}14"] = f"=-MIN({col}12+{col}13,MAX({col}10+{col}13,0))"
        m[f"{col}15"] = f"={col}12+{col}13+{col}14"
        m[f"{col}16"] = f"=-{col}12*(VLOOKUP({col}$2,SOFR!$A$2:$B$6,2,FALSE)+Inputs!$B$15)"
        m[f"{col}18"] = f"={prev}18+{col}10+{col}13+{col}14" if prev else f"={col}10+{col}13+{col}14"
        m[f"{col}19"] = f"={col}15-{col}18"
        m[f"{col}21"] = f"={col}10+{col}13+{col}14-({col}18-{prev}18)" if prev else f"={col}10+{col}13+{col}14-{col}18"
    m["F24"] = "=F5*8.5-F15+F18" if hidden_hardcode else "=F5*Inputs!$B$17-F15+F18"
    if hidden_hardcode:
        m["F27"] = "=F26*1.075"           # second untraceable constant
        m["F28"] = "=F27*1.075"
        m["H2"] = '=BDP("SPX Index","PX_LAST")'  # data-vendor formula
    m["F25"] = "=F24/Inputs!$B$18"
    m["F26"] = "=(F24/Inputs!$B$18)^(1/4)-1"
    m["F26"].number_format = "0.0%"
    wb.save(path)


def build_input(path: Path, hints: bool = False):
    """Input = gold with Model rows below revenue/EBITDA removed and cases hardcoded."""
    build_gold(path)
    wb = openpyxl.load_workbook(path)
    inp = wb["Inputs"]
    for coord in ("B9", "B10", "B11", "B12", "C4", "D4", "C5", "D5", "C6", "D6", "C7", "D7", "B14", "B15", "B16", "B17", "B18"):
        inp[coord] = None
    for coord in ("A9", "A10", "A11", "A12", "A14", "A15", "A16", "A17", "A18"):
        inp[coord] = None
    inp["B1"] = None
    inp["A1"] = None
    m = wb["Model"]
    for row in range(7, 27):
        for col in "ABCDEF":
            m[f"{col}{row}"] = None
    for col in "CDEF":
        m[f"{col}3"] = f"={chr(ord(col)-1)}3*(1+0.10)"
        m[f"{col}5"] = f"={col}3*0.30"
    if hints:
        # illogical build: pre-populated downstream values equal to the gold's formula results
        m["A10"] = "Free Cash Flow Available for Debt Service"
        m["F10"] = 329.0
        m["A19"] = "Net Debt"
        m["F19"] = -984.3
    wb.save(path)


def build_bundle(root: Path, good: bool = True, hints: bool = False, audio_bytes: int = 200_000) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    build_input(root / "Meridian-abc123-input.xlsx", hints=hints)
    build_gold(root / "Meridian-abc123-gold-output.xlsx", hidden_hardcode=not good)
    (root / "rubric.md").write_text(RUBRIC_GOOD if good else RUBRIC_BAD)
    (root / "script.txt").write_text(SCRIPT)
    (root / "MD Call Recording.m4a").write_bytes(b"\0" * audio_bytes)
    return root


if __name__ == "__main__":
    import sys
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "fixture_task")
    build_bundle(out / "good", good=True)
    build_bundle(out / "bad", good=False, hints=True)
    print("built", out)
